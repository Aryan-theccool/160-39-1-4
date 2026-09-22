"""Unit and integration tests for unified multi-format document reader (PDF, DOCX, TXT)."""

import io
import pytest
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from bis_api.document_reader import (
    DEFAULT_MAX_BYTES,
    DEFAULT_MAX_CHARS,
    DocumentContent,
    DocumentPage,
    DocumentParagraph,
    DocumentReadResult,
    DocumentSource,
    DocumentTable,
    detect_format,
    read_document,
    read_docx,
    read_pdf,
    read_txt,
)


def _create_synthetic_pdf(page_texts: list[str]) -> bytes:
    """Create an in-memory PDF with specified text per page using reportlab."""
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    for text in page_texts:
        if text:
            c.drawString(50, 700, text)
        else:
            # draw a non-text rectangle for empty/scanned simulation
            c.rect(50, 50, 100, 100, stroke=1, fill=0)
        c.showPage()
    c.save()
    return buf.getvalue()


def _create_synthetic_docx() -> bytes:
    """Create an in-memory DOCX with headings, paragraphs, list items, and tables."""
    import docx

    doc = docx.Document()
    doc.add_heading("Technical Specification for LED Luminaire", level=1)
    doc.add_paragraph("The contractor shall furnish and install outdoor street lighting.")
    doc.add_paragraph("Driver efficiency shall exceed 90 percent.", style="List Bullet")

    table = doc.add_table(rows=3, cols=3)
    table.rows[0].cells[0].text = "Parameter"
    table.rows[0].cells[1].text = "Requirement"
    table.rows[0].cells[2].text = "Standard"

    table.rows[1].cells[0].text = "Power"
    table.rows[1].cells[1].text = "90 W"
    table.rows[1].cells[2].text = "IS 10322 (Part 5/Sec 3)"

    table.rows[2].cells[0].text = "Ingress Protection"
    table.rows[2].cells[1].text = "IP66"
    table.rows[2].cells[2].text = "IS/IEC 60529"

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# TEST 1: PDF Multi-page Extraction
# ---------------------------------------------------------------------------
def test_pdf_multi_page_extraction():
    """Verify multi-page PDF preserves page boundaries, page count, and full text."""
    p1 = "Page 1: Scope of procurement for municipal outdoor illumination."
    p2 = "Page 2: Mandatory adherence to IS 10322 Part 5 Section 3."
    pdf_bytes = _create_synthetic_pdf([p1, p2])

    result = read_document(pdf_bytes, filename="tender_spec.pdf")

    assert result.ok is True
    assert result.document is not None
    doc = result.document

    assert doc.source.source_type == "pdf"
    assert doc.source.extension == ".pdf"
    assert doc.page_count == 2
    assert len(doc.pages) == 2

    assert doc.pages[0].page_number == 1
    assert "Page 1: Scope of procurement" in doc.pages[0].text
    assert doc.pages[0].character_count == len(doc.pages[0].text)

    assert doc.pages[1].page_number == 2
    assert "Page 2: Mandatory adherence" in doc.pages[1].text
    assert doc.pages[1].character_count == len(doc.pages[1].text)

    assert "Scope of procurement" in doc.full_text
    assert "Mandatory adherence" in doc.full_text
    assert doc.character_count == len(doc.full_text)
    assert doc.extraction_method in ("pymupdf", "pypdf", "builtin")


# ---------------------------------------------------------------------------
# TEST 2: DOCX Extraction
# ---------------------------------------------------------------------------
def test_docx_extraction_structure():
    """Verify DOCX extracts headings, paragraphs, bullet lists, and tables structurally."""
    docx_bytes = _create_synthetic_docx()

    result = read_document(docx_bytes, filename="tender_scope.docx")

    assert result.ok is True
    assert result.document is not None
    doc = result.document

    assert doc.source.source_type == "docx"
    assert doc.source.extension == ".docx"
    assert doc.extraction_method == "python-docx"

    # Verify paragraphs & styles
    styles = [p.style for p in doc.paragraphs]
    texts = [p.text for p in doc.paragraphs]

    assert any("Heading" in s for s in styles)
    assert any("Technical Specification for LED Luminaire" in t for t in texts)
    assert any("outdoor street lighting" in t for t in texts)
    assert any("Driver efficiency" in t for t in texts)

    # Verify tables
    assert len(doc.tables) == 1
    tbl = doc.tables[0]
    assert tbl.table_index == 0
    assert len(tbl.rows) == 3
    assert tbl.rows[0] == ["Parameter", "Requirement", "Standard"]
    assert tbl.rows[1] == ["Power", "90 W", "IS 10322 (Part 5/Sec 3)"]
    assert tbl.rows[2] == ["Ingress Protection", "IP66", "IS/IEC 60529"]

    # Verify full_text incorporates table rows
    assert "IP66" in doc.full_text
    assert "90 W" in doc.full_text
    assert doc.character_count == len(doc.full_text)


# ---------------------------------------------------------------------------
# TEST 3: TXT UTF-8 & Legacy Encoding
# ---------------------------------------------------------------------------
def test_txt_utf8_extraction():
    """Verify plain text extraction preserves text and line boundaries."""
    raw_text = (
        "GOVERNMENT TENDER NOTICE 2026\n"
        "Clause 1: Supply of Galvanized Steel Pipes.\n"
        "Clause 2: Product must comply with IS 1239 (Part 1).\n"
        "Clause 3: Minimum tensile strength 320 MPa.\n"
    )
    txt_bytes = raw_text.encode("utf-8")

    result = read_document(txt_bytes, filename="requirements.txt")

    assert result.ok is True
    assert result.document is not None
    doc = result.document

    assert doc.source.source_type == "txt"
    assert doc.full_text == raw_text.strip()
    assert doc.character_count == len(doc.full_text)
    assert len(doc.paragraphs) == 4
    assert doc.paragraphs[1].text == "Clause 1: Supply of Galvanized Steel Pipes."


def test_txt_latin1_fallback():
    """Verify Latin-1 encoded text file decodes successfully with fallback."""
    # Byte with high-bit Latin-1 character e.g. 'é' or degree symbol '°'
    raw_bytes = "Operating temperature range: -10\xb0C to +50\xb0C.\nIS 10322 compliance mandatory.".encode("latin-1")

    result = read_document(raw_bytes, filename="spec_latin1.txt")

    assert result.ok is True
    assert result.document is not None
    assert "Operating temperature" in result.document.full_text
    assert "IS 10322" in result.document.full_text


# ---------------------------------------------------------------------------
# TEST 4: Empty Files
# ---------------------------------------------------------------------------
def test_empty_txt_rejection():
    """Verify empty TXT yields clean failure."""
    result = read_document(b"", filename="empty.txt")
    assert result.ok is False
    assert result.error_code == "EMPTY_FILE"

    result_whitespace = read_document(b"   \n\n\t  ", filename="blank.txt")
    assert result_whitespace.ok is False
    assert result_whitespace.error_code == "EMPTY_FILE"


def test_empty_pdf_rejection():
    """Verify empty PDF bytes yield clean failure."""
    result = read_document(b"", filename="empty.pdf")
    assert result.ok is False
    assert result.error_code == "EMPTY_FILE"


# ---------------------------------------------------------------------------
# TEST 5: Malformed Documents
# ---------------------------------------------------------------------------
def test_malformed_pdf_rejection():
    """Verify non-PDF file with .pdf extension fails cleanly."""
    result = read_document(b"This is not a real PDF document at all.", filename="corrupt.pdf")
    assert result.ok is False
    assert result.error_code in ("MALFORMED_DOCUMENT", "CORRUPT_FILE")


def test_malformed_docx_rejection():
    """Verify non-docx data fails cleanly."""
    result = read_document(b"PK fake zip content that is not docx", filename="bad.docx")
    assert result.ok is False
    assert result.error_code == "MALFORMED_DOCUMENT"


# ---------------------------------------------------------------------------
# TEST 6: Scanned / Image-Only PDF (No OCR)
# ---------------------------------------------------------------------------
def test_scanned_pdf_no_ocr_rejection():
    """Verify image-only / blank PDF fails with NO_TEXT_LAYER and does NOT run OCR."""
    pdf_bytes = _create_synthetic_pdf(["", ""])  # Two blank/shape-only pages

    result = read_document(pdf_bytes, filename="scanned_tender.pdf")

    assert result.ok is False
    assert result.error_code == "NO_TEXT_LAYER"
    assert "scanned or image-only" in result.error_message or "usable text layer" in result.error_message
    assert "OCR is not attempted" in result.error_message


# ---------------------------------------------------------------------------
# TEST 7: Unsupported Extension
# ---------------------------------------------------------------------------
def test_unsupported_extension_rejection():
    """Verify unsupported file format (e.g. .exe, .zip) is cleanly rejected."""
    result = read_document(b"MZ\x90\x00\x03\x00\x00\x00", filename="trojan.exe")
    assert result.ok is False
    assert result.error_code == "UNSUPPORTED_FORMAT"
    assert "Unsupported document format" in result.error_message


# ---------------------------------------------------------------------------
# TEST 8: File Size Limit Validation
# ---------------------------------------------------------------------------
def test_file_size_validation():
    """Verify oversized file is rejected prior to parsing."""
    small_limit = 50  # 50 bytes max
    data = b"A" * 100

    result = read_document(data, filename="large.txt", max_bytes=small_limit)
    assert result.ok is False
    assert result.error_code == "FILE_TOO_LARGE"
    assert "exceeds the maximum allowed limit" in result.error_message


# ---------------------------------------------------------------------------
# TEST 9: Character Counts & Serialization
# ---------------------------------------------------------------------------
def test_character_count_and_dict_serialization():
    """Verify character count is deterministic and as_dict serialization works."""
    raw_text = "Standard Specification Requirement: Minimum thickness 2.5 mm."
    result = read_document(raw_text.encode("utf-8"), filename="test.txt")

    assert result.ok is True
    doc = result.document
    assert doc is not None
    assert doc.character_count == len(raw_text)

    d = result.as_dict()
    assert d["ok"] is True
    assert d["document"]["source"]["filename"] == "test.txt"
    assert d["document"]["character_count"] == len(raw_text)
    assert len(d["document"]["paragraphs"]) == 1
    assert d["document"]["paragraphs"][0]["text"] == raw_text


# ---------------------------------------------------------------------------
# TEST 10: Format Detection
# ---------------------------------------------------------------------------
def test_format_detection():
    """Verify format detection handles filenames, MIME types, and magic bytes."""
    assert detect_format(b"%PDF-1.4...", filename="") == "pdf"
    assert detect_format(b"PK\x03\x04...", filename="") == "docx"
    assert detect_format(b"", filename="doc.pdf") == "pdf"
    assert detect_format(b"", filename="doc.docx") == "docx"
    assert detect_format(b"", filename="doc.txt") == "txt"
    assert detect_format(b"", media_type="application/pdf") == "pdf"
    assert detect_format(b"", media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document") == "docx"
    assert detect_format(b"", media_type="text/plain") == "txt"
    assert detect_format(b"", filename="unknown.xyz") == "unknown"
