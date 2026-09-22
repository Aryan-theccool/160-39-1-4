"""Tests for professional downloadable procurement report generation (DOCX and PDF)."""

from __future__ import annotations

import io
from pathlib import Path
import pytest
import docx
from pypdf import PdfReader
from fastapi.testclient import TestClient

from bis_api.config import Settings
from bis_api.main import create_app
from bis_rag.report_builder import build_report_docx, build_report_pdf, normalize_report_data


@pytest.fixture
def client():
    settings = Settings(
        data_dir=Path("."),
        embedder="hash",
        warm_indexes=False,
        llm="extractive",
    )
    app = create_app(settings)
    with TestClient(app) as test_client:
        yield test_client


@pytest.fixture
def full_report_data():
    return {
        "question": "Outdoor LED street light for municipal road lighting, approximately 90 W, IP66, suitable for outdoor use.",
        "answer": "Primary standard for street lighting luminaires is IS 10322 (Part 5 / Sec 3).",
        "mode": "llm",
        "generator": "groq",
        "grounded": True,
        "report": {
            "summary": "This procurement assessment establishes the applicable Indian Standards for outdoor LED street lighting equipment.",
            "product_understanding": {
                "product": "Outdoor LED Street Light Luminaire",
                "application": "Municipal & Highway Illumination",
                "specifications": {
                    "Wattage": "90W",
                    "Ingress Protection": "IP66",
                    "CCT": "4000K",
                    "Efficacy": ">= 120 lm/W"
                }
            },
            "applicable_standards": [
                {
                    "designation": "IS 10322 (Part 5 / Sec 3):2012",
                    "title": "Luminaires - Part 5 Particular Requirements - Sec 3 Luminaires for Road and Street Lighting",
                    "status": "current",
                    "relevance": "Primary Standard",
                    "reason": "Direct standard governing street lighting luminaire safety, thermal performance, and mechanical construction."
                },
                {
                    "designation": "IS 16102 (Part 1):2014",
                    "title": "Self-Ballasted LED Lamps for General Lighting Services - Part 1 Safety Requirements",
                    "status": "current",
                    "relevance": "Applicable Standard",
                    "reason": "Safety requirements for self-ballasted LED light sources."
                }
            ],
            "related_standards": [
                {
                    "designation": "IS 16103 (Part 1):2012",
                    "title": "Led Modules for General Lighting Part 1 Safety Requirements",
                    "relationship": "Component Standard"
                }
            ],
            "compliance_notes": [
                "Mandatory BIS CRS registration applies to LED luminaires under the Electronics and IT Goods QCO.",
                "Verify safety testing compliance reports against IS 10322."
            ],
            "warnings": [
                "Ensure luminaire mounting bracket meets IS 875 wind-load standards."
            ],
            "evidence": [
                {
                    "designation": "IS 10322 (Part 5 / Sec 3):2012",
                    "chunk_id": "chunk_001",
                    "reason": "Specifies road lighting luminaire insulation and IP protection."
                }
            ]
        },
        "citations": [
            {
                "designation": "IS 10322 (Part 5 / Sec 3):2012",
                "title": "Luminaires for Road and Street Lighting",
                "is_current": True,
                "archive_url": "https://archive.org/details/gov.in.is.10322.5.3.2012"
            }
        ],
        "unsupported_citations": [],
        "warnings": [
            "Ensure luminaire mounting bracket meets IS 875 wind-load standards."
        ]
    }


def test_docx_generation_structure_and_validity(full_report_data):
    """Test 1: DOCX generation produces a valid, readable .docx with expected content."""
    docx_buf = build_report_docx(full_report_data)
    assert isinstance(docx_buf, io.BytesIO)
    content = docx_buf.getvalue()
    assert len(content) > 1000

    # Load with python-docx to verify binary validity
    doc = docx.Document(docx_buf)
    all_text = " ".join(p.text for p in doc.paragraphs)
    
    # Check key sections are present
    assert "AI-Powered Indian Standards Recommendation Report" in all_text
    assert "Executive Summary" in all_text
    assert "Product & Procurement Understanding" in all_text
    assert "Applicable Indian Standards" in all_text
    assert "Related & Allied Standards" in all_text
    assert "Compliance & Quality Notes" in all_text
    assert "Important Limitations" in all_text
    
    # Check tables exist
    assert len(doc.tables) >= 2  # metadata table + applicable standards table
    
    # Check standard numbers preserved exactly
    table_text = " ".join(cell.text for t in doc.tables for row in t.rows for cell in row.cells)
    assert "IS 10322 (Part 5 / Sec 3):2012" in table_text
    assert "IS 16102 (Part 1):2014" in table_text


def test_pdf_generation_structure_and_validity(full_report_data):
    """Test 2: PDF generation produces a valid, readable PDF with expected text."""
    pdf_buf = build_report_pdf(full_report_data)
    assert isinstance(pdf_buf, io.BytesIO)
    content = pdf_buf.getvalue()
    assert len(content) > 1000
    assert content.startswith(b"%PDF")

    # Read with PdfReader
    reader = PdfReader(io.BytesIO(content))
    assert len(reader.pages) >= 1
    extracted_text = " ".join(page.extract_text() for page in reader.pages)
    
    assert "AI-Powered Indian Standards Recommendation Report" in extracted_text
    assert "Executive Summary" in extracted_text
    assert "Applicable Indian Standards" in extracted_text
    assert "IS 10322" in extracted_text


def test_api_docx_export_endpoint(client, full_report_data):
    """Test 3: POST /api/v1/report/docx returns 200, attachment header, and valid binary."""
    resp = client.post("/api/v1/report/docx", json=full_report_data)
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert "attachment" in resp.headers["content-disposition"]
    assert "bis_standards_report_" in resp.headers["content-disposition"]
    assert resp.headers["content-disposition"].endswith('.docx"')

    # Verify response body parses as docx
    doc = docx.Document(io.BytesIO(resp.content))
    assert len(doc.paragraphs) > 5


def test_api_pdf_export_endpoint(client, full_report_data):
    """Test 4: POST /api/v1/report/pdf returns 200, attachment header, and valid PDF."""
    resp = client.post("/api/v1/report/pdf", json=full_report_data)
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/pdf"
    assert "attachment" in resp.headers["content-disposition"]
    assert "bis_standards_report_" in resp.headers["content-disposition"]
    assert resp.headers["content-disposition"].endswith('.pdf"')

    # Verify response body starts with %PDF
    assert resp.content.startswith(b"%PDF")
    reader = PdfReader(io.BytesIO(resp.content))
    assert len(reader.pages) >= 1
    extracted_text = " ".join(page.extract_text() for page in reader.pages)
    assert "AI-Powered Indian Standards Recommendation Report" in extracted_text


def test_missing_optional_fields_resilience():
    """Test 5: Report builder handles sparse or empty optional fields safely."""
    minimal_data = {
        "question": "General steel requirement",
        "answer": "See IS 2062 for structural steel.",
        "report": None,
        "citations": []
    }
    
    docx_buf = build_report_docx(minimal_data)
    pdf_buf = build_report_pdf(minimal_data)
    
    assert len(docx_buf.getvalue()) > 500
    assert len(pdf_buf.getvalue()) > 500


def test_grounding_warnings_and_unsupported_citations_preserved():
    """Test 6: Unsupported citations and supersession warnings are clearly rendered."""
    unsupported_data = {
        "question": "Drinking water filter",
        "answer": "Water filter standard answer.",
        "grounded": False,
        "warnings": [
            "IS 456:1978 is superseded by IS 456:2000."
        ],
        "unsupported_citations": ["IS 10500", "IS 1070"],
        "report": {
            "summary": "Water quality screening.",
            "applicable_standards": [],
            "warnings": ["IS 456:1978 is superseded by IS 456:2000."]
        },
        "citations": []
    }

    docx_buf = build_report_docx(unsupported_data)
    doc = docx.Document(docx_buf)
    all_text = " ".join(p.text for p in doc.paragraphs)
    table_text = " ".join(cell.text for t in doc.tables for row in t.rows for cell in row.cells)
    combined = all_text + " " + table_text

    assert "IS 10500" in combined
    assert "IS 1070" in combined
    assert "superseded" in combined.lower()

    pdf_buf = build_report_pdf(unsupported_data)
    reader = PdfReader(io.BytesIO(pdf_buf.getvalue()))
    pdf_text = " ".join(page.extract_text() for page in reader.pages)
    assert "IS 10500" in pdf_text
    assert "IS 1070" in pdf_text


def test_long_content_overflow_handling():
    """Test 7: Handles long titles, multiple standards and long text without error."""
    long_data = {
        "question": "Extremely detailed multi-parameter query " * 10,
        "answer": "Long answer body text " * 50,
        "report": {
            "summary": "Detailed executive summary paragraph " * 15,
            "product_understanding": {
                "product": "Complex Industrial Assembly " * 3,
                "application": "High Voltage Switchgear Station",
                "specifications": {f"Spec_{i}": f"Value parameter specification {i}" * 2 for i in range(10)}
            },
            "applicable_standards": [
                {
                    "designation": f"IS {1000 + i}:2020",
                    "title": f"Title for Indian Standard {1000 + i} Specification and Requirements " * 3,
                    "status": "current" if i % 2 == 0 else "superseded",
                    "relevance": "Primary Standard",
                    "reason": f"Detailed rationale for standard {1000 + i} applicability in this procurement. " * 3
                }
                for i in range(8)
            ],
            "related_standards": [
                {
                    "designation": f"IS {2000 + i}:2018",
                    "title": f"Allied Reference Standard {2000 + i} " * 2,
                    "relationship": "Test Method"
                }
                for i in range(5)
            ],
            "compliance_notes": [f"Compliance note requirement {i} with regulatory context." for i in range(6)],
            "warnings": [f"Crucial warning note {i} regarding installation compliance." for i in range(4)]
        },
        "citations": [
            {
                "designation": f"IS {1000 + i}:2020",
                "title": f"Title for Indian Standard {1000 + i}",
                "archive_url": f"https://archive.org/details/gov.in.is.{1000 + i}.2020"
            }
            for i in range(8)
        ]
    }

    docx_buf = build_report_docx(long_data)
    doc = docx.Document(docx_buf)
    assert len(doc.paragraphs) > 20

    pdf_buf = build_report_pdf(long_data)
    reader = PdfReader(io.BytesIO(pdf_buf.getvalue()))
    assert len(reader.pages) >= 2  # Multipage document generated cleanly


