"""Unit and integration tests for Tender Clause & Technical Requirement Extractor."""

import io
import pytest
import docx
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from bis_api.document_reader import (
    DocumentContent,
    DocumentParagraph,
    DocumentSource,
    DocumentTable,
    read_document,
)
from bis_rag.tender_extractor import (
    TenderClause,
    TenderExtractionResult,
    TenderRequirement,
    classify_clause_type,
    extract_from_document,
    extract_requirements_from_table,
    extract_requirements_from_text,
    extract_tender_requirements,
    find_standard_references,
    normalize_parameter_name,
    parse_value_and_unit,
)


def _create_sample_docx_tender() -> bytes:
    """Create an in-memory DOCX tender document with clauses, tables, and standard citations."""
    doc = docx.Document()
    doc.add_heading("Section 1.0 Scope of Procurement", level=1)
    doc.add_paragraph("The scope covers the supply and installation of outdoor LED street light luminaires.")

    doc.add_heading("Section 2.0 Technical Specifications", level=1)
    doc.add_paragraph("All luminaires shall satisfy the following engineering parameters:")

    table = doc.add_table(rows=6, cols=4)
    table.rows[0].cells[0].text = "Parameter"
    table.rows[0].cells[1].text = "Specified Requirement"
    table.rows[0].cells[2].text = "Unit"
    table.rows[0].cells[3].text = "Applicable Standard"

    table.rows[1].cells[0].text = "Rated Power"
    table.rows[1].cells[1].text = "90"
    table.rows[1].cells[2].text = "W"
    table.rows[1].cells[3].text = "IS 10322 (Part 5 / Section 3)"

    table.rows[2].cells[0].text = "Luminous Efficacy"
    table.rows[2].cells[1].text = ">= 120"
    table.rows[2].cells[2].text = "lm/W"
    table.rows[2].cells[3].text = "IS 16107 (Part 2 / Section 1)"

    table.rows[3].cells[0].text = "Ingress Protection"
    table.rows[3].cells[1].text = "IP66"
    table.rows[3].cells[2].text = ""
    table.rows[3].cells[3].text = "IS/IEC 60529"

    table.rows[4].cells[0].text = "Surge Protection"
    table.rows[4].cells[1].text = ">= 10"
    table.rows[4].cells[2].text = "kV"
    table.rows[4].cells[3].text = "IS 15885 (Part 2 / Section 13)"

    table.rows[5].cells[0].text = "Color Temperature"
    table.rows[5].cells[1].text = "4000 to 5000"
    table.rows[5].cells[2].text = "K"
    table.rows[5].cells[3].text = "IS 16102 (Part 1)"

    doc.add_heading("Section 3.0 Warranty and Commercial Terms", level=1)
    doc.add_paragraph("The bidder shall provide a comprehensive warranty of minimum 5 years from date of commissioning.")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# TEST 1: Value & Unit Parsing
# ---------------------------------------------------------------------------
def test_parse_value_and_unit():
    """Verify numeric, range, operator, and unit parsing."""
    num, num_max, unit, op = parse_value_and_unit("90 W")
    assert num == 90.0
    assert num_max is None
    assert unit == "W"
    assert op == "exact"

    num, num_max, unit, op = parse_value_and_unit(">= 120 lm/W")
    assert num == 120.0
    assert unit == "lm/W"
    assert op == ">="

    num, num_max, unit, op = parse_value_and_unit("4000 to 5000 K")
    assert num == 4000.0
    assert num_max == 5000.0
    assert unit == "K"
    assert op == "range"

    num, num_max, unit, op = parse_value_and_unit("IP66")
    assert unit == "IP"
    assert op == "exact"


# ---------------------------------------------------------------------------
# TEST 2: Parameter Normalization
# ---------------------------------------------------------------------------
def test_normalize_parameter_name():
    """Verify canonical naming and categorization for common tender parameters."""
    canon, cat = normalize_parameter_name("rated wattage")
    assert canon == "Power"
    assert cat == "electrical"

    canon, cat = normalize_parameter_name("system efficacy (lm/w)")
    assert canon == "Luminous Efficacy"
    assert cat == "photometric"

    canon, cat = normalize_parameter_name("degree of protection (IP rating)")
    assert canon == "Ingress Protection (IP)"
    assert cat == "environmental"

    canon, cat = normalize_parameter_name("tensile strength (MPa)")
    assert canon == "Tensile Strength"
    assert cat == "materials"

    canon, cat = normalize_parameter_name("comprehensive guarantee")
    assert canon == "Warranty"
    assert cat == "commercial"


# ---------------------------------------------------------------------------
# TEST 3: Standard Reference Detection
# ---------------------------------------------------------------------------
def test_find_standard_references():
    """Verify standard citations are extracted and canonicalized."""
    text = (
        "Luminaires must conform to IS 10322 (Part 5/Sec 3):2012, "
        "IS 16107 (Part 2/Section 1): 2012, IS/IEC 60529, and ISO 9001 quality management."
    )
    stds = find_standard_references(text)
    assert any("IS 10322" in s for s in stds)
    assert any("IS 16107" in s for s in stds)
    assert any("60529" in s for s in stds)
    assert any("ISO 9001" in s for s in stds)


# ---------------------------------------------------------------------------
# TEST 4: Clause Classification
# ---------------------------------------------------------------------------
def test_classify_clause_type():
    """Verify classification of clause types."""
    assert classify_clause_type("Electrical ratings and photometric performance", "Technical Specifications") == "technical"
    assert classify_clause_type("The luminaires shall comply with BIS standards and IS codes", "Standards & Codes") == "standards"
    assert classify_clause_type("Warranty period shall be 5 years. EMD is required.", "Commercial Terms") == "commercial"
    assert classify_clause_type("Supply, testing, and commissioning of street lights", "Scope of Work") == "scope"
    assert classify_clause_type("Bidder must have minimum 3 years experience and 10 Cr turnover", "Eligibility Criteria") == "eligibility"


# ---------------------------------------------------------------------------
# TEST 5: Table Requirement Extraction
# ---------------------------------------------------------------------------
def test_extract_requirements_from_table():
    """Verify structured requirement extraction from a DocumentTable."""
    tbl = DocumentTable(
        table_index=0,
        rows=[
            ["Parameter", "Requirement", "Unit", "Standard"],
            ["Power Consumption", "<= 90", "W", "IS 10322"],
            ["Luminous Efficacy", ">= 120", "lm/W", "IS 16107"],
            ["Ingress Protection", "IP66", "", "IS/IEC 60529"],
        ],
    )

    reqs = extract_requirements_from_table(tbl)
    assert len(reqs) == 3

    assert reqs[0].parameter == "Power"
    assert reqs[0].numeric_value == 90.0
    assert reqs[0].operator == "<="
    assert reqs[0].unit == "W"
    assert any("IS 10322" in s for s in reqs[0].standard_references)

    assert reqs[1].parameter == "Luminous Efficacy"
    assert reqs[1].numeric_value == 120.0
    assert reqs[1].operator == ">="
    assert reqs[1].unit == "lm/W"

    assert reqs[2].parameter == "Ingress Protection (IP)"
    assert reqs[2].category == "environmental"
    assert any("60529" in s for s in reqs[2].standard_references)


# ---------------------------------------------------------------------------
# TEST 6: Free Text Requirement Extraction
# ---------------------------------------------------------------------------
def test_extract_requirements_from_text():
    """Verify regex-based requirement extraction from free-text clauses."""
    text = (
        "The LED street luminaire shall have rated power of 90 W, luminous efficacy not less than 120 lm/W, "
        "color temperature 4000 K to 5000 K, CRI >= 70, power factor >= 0.95, THD <= 10%, "
        "surge protection of 10 kV, ingress protection minimum IP66, and warranty of 5 years. "
        "Standard compliance: IS 10322 (Part 5/Sec 3)."
    )

    reqs = extract_requirements_from_text(text, source_location="Paragraph 2")
    assert len(reqs) >= 8

    params = {r.parameter: r for r in reqs}
    assert "Power" in params
    assert params["Power"].numeric_value == 90.0

    assert "Luminous Efficacy" in params
    assert params["Luminous Efficacy"].numeric_value == 120.0

    assert "Correlated Color Temperature (CCT)" in params
    assert params["Correlated Color Temperature (CCT)"].numeric_value == 4000.0
    assert params["Correlated Color Temperature (CCT)"].numeric_max == 5000.0

    assert "Color Rendering Index (CRI)" in params
    assert params["Color Rendering Index (CRI)"].numeric_value == 70.0

    assert "Power Factor" in params
    assert params["Power Factor"].numeric_value == 0.95

    assert "Total Harmonic Distortion (THD)" in params
    assert params["Total Harmonic Distortion (THD)"].numeric_value == 10.0

    assert "Surge Protection" in params
    assert params["Surge Protection"].numeric_value == 10.0

    assert "Ingress Protection (IP)" in params
    assert "IP66" in params["Ingress Protection (IP)"].specified_value

    assert "Warranty" in params
    assert params["Warranty"].numeric_value == 5.0


# ---------------------------------------------------------------------------
# TEST 7: End-to-End Extraction from DOCX Document
# ---------------------------------------------------------------------------
def test_extract_from_docx_document():
    """Verify end-to-end requirement extraction from a generated DOCX tender."""
    docx_bytes = _create_sample_docx_tender()

    result = extract_from_document(docx_bytes, filename="tender_rfp.docx")

    assert result.total_requirements >= 5
    assert len(result.clauses) >= 3
    assert any(c.clause_type == "technical" for c in result.clauses)
    assert any(c.clause_type == "commercial" for c in result.clauses)

    stds = result.detected_standards
    assert any("IS 10322" in s for s in stds)
    assert any("IS 16107" in s for s in stds)
    assert any("60529" in s for s in stds)

    # Verify serialization
    d = result.as_dict()
    assert d["total_requirements"] == result.total_requirements
    assert len(d["requirements"]) == len(result.requirements)
    assert d["requirements"][0]["parameter"] != ""


# ---------------------------------------------------------------------------
# TEST 8: End-to-End Extraction from TXT Document
# ---------------------------------------------------------------------------
def test_extract_from_txt_document():
    """Verify end-to-end requirement extraction from a plain text tender."""
    txt_content = (
        "TENDER FOR GALVANIZED STEEL PIPES\n\n"
        "1.0 Technical Specification:\n"
        "Pipes must comply with IS 1239 (Part 1). Tensile strength shall be minimum 320 MPa.\n"
        "Yield strength shall be minimum 250 MPa. Warranty of 2 years required.\n"
    )

    result = extract_from_document(txt_content.encode("utf-8"), filename="steel_tender.txt")

    assert result.total_requirements >= 2
    assert any("IS 1239" in s for s in result.detected_standards)
    params = [r.parameter for r in result.requirements]
    assert "Tensile Strength" in params
    assert "Yield Strength" in params
