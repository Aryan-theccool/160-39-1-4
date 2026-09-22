"""Unit and integration tests for Tender API Endpoints and Report Export (PDF/DOCX)."""

import io
from pathlib import Path
import pytest
import docx
from fastapi.testclient import TestClient
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from bis_api.config import Settings
from bis_api.main import create_app
from synth_bis_data import make_synth_data


@pytest.fixture(scope="module")
def app_instance(tmp_path_factory):
    data_path = make_synth_data(tmp_path_factory.mktemp("bis_data"))
    settings = Settings(
        data_dir=data_path,
        embedder="hash",
        warm_indexes=False,
        llm="extractive",
    )
    return create_app(settings)


@pytest.fixture
def client(app_instance):
    with TestClient(app_instance) as test_client:
        yield test_client


def _create_sample_pdf() -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    c.drawString(50, 700, "Scope: Supply of LED street lighting systems.")
    c.drawString(50, 680, "Technical Requirements: Minimum power 90 W, efficacy >= 120 lm/W, IP rating IP66.")
    c.drawString(50, 660, "Mandatory compliance with IS 10322 (Part 5 / Sec 3) and IS 16107 (Part 2 / Sec 1).")
    c.showPage()
    c.save()
    return buf.getvalue()


def _create_sample_docx() -> bytes:
    doc = docx.Document()
    doc.add_heading("Technical Specification for Street Lighting", level=1)
    doc.add_paragraph("Supply of energy-efficient LED luminaires.")
    table = doc.add_table(rows=3, cols=3)
    table.rows[0].cells[0].text = "Parameter"
    table.rows[0].cells[1].text = "Requirement"
    table.rows[0].cells[2].text = "Standard"

    table.rows[1].cells[0].text = "Rated Power"
    table.rows[1].cells[1].text = "90 W"
    table.rows[1].cells[2].text = "IS 10322 (Part 5/Sec 3)"

    table.rows[2].cells[0].text = "Ingress Protection"
    table.rows[2].cells[1].text = "IP66"
    table.rows[2].cells[2].text = "IS/IEC 60529"

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# TEST 1: PDF File Upload Analysis
# ---------------------------------------------------------------------------
def test_tender_analyze_file_pdf(client):
    """Verify PDF upload to /api/v1/tender/analyze/file returns structured compliance analysis."""
    pdf_bytes = _create_sample_pdf()
    files = {"file": ("tender_spec.pdf", pdf_bytes, "application/pdf")}
    data = {"product_description": "LED street lighting luminaire", "sector": "electrical"}

    response = client.post("/api/v1/tender/analyze/file", files=files, data=data)
    assert response.status_code == 200

    body = response.json()
    assert body["filename"] == "tender_spec.pdf"
    assert body["source_type"] == "pdf"
    assert len(body["requirements"]) >= 2
    assert any("IS 10322" in s for s in body["detected_standards"])
    assert "compliance_matrix" in body
    assert body["compliance_matrix"]["compliance_score"] > 0.5


# ---------------------------------------------------------------------------
# TEST 2: DOCX File Upload Analysis
# ---------------------------------------------------------------------------
def test_tender_analyze_file_docx(client):
    """Verify DOCX upload to /api/v1/tender/analyze/file parses tables and evaluates compliance."""
    docx_bytes = _create_sample_docx()
    files = {"file": ("tender_spec.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}

    response = client.post("/api/v1/tender/analyze/file", files=files)
    assert response.status_code == 200

    body = response.json()
    assert body["filename"] == "tender_spec.docx"
    assert body["source_type"] == "docx"
    assert len(body["requirements"]) >= 2
    assert "compliance_matrix" in body
    assert len(body["compliance_matrix"]["evaluations"]) >= 2


# ---------------------------------------------------------------------------
# TEST 3: TXT File Upload Analysis
# ---------------------------------------------------------------------------
def test_tender_analyze_file_txt(client):
    """Verify TXT upload to /api/v1/tender/analyze/file processes plain text RFP."""
    txt_content = (
        "Technical Specifications:\n"
        "1. Pipes shall conform to IS 1239 (Part 1).\n"
        "2. Minimum tensile strength: 320 MPa.\n"
        "3. Minimum yield strength: 250 MPa.\n"
    )
    files = {"file": ("tender_spec.txt", txt_content.encode("utf-8"), "text/plain")}

    response = client.post("/api/v1/tender/analyze/file", files=files)
    assert response.status_code == 200

    body = response.json()
    assert body["filename"] == "tender_spec.txt"
    assert body["source_type"] == "txt"
    assert any("Tensile Strength" in r["parameter"] for r in body["requirements"])


# ---------------------------------------------------------------------------
# TEST 4: Direct Text Analysis Endpoint
# ---------------------------------------------------------------------------
def test_tender_analyze_text_endpoint(client):
    """Verify direct JSON text submission to /api/v1/tender/analyze/text."""
    payload = {
        "text": (
            "Section 1: The outdoor luminaire shall have rated wattage of 90 W, "
            "ingress protection minimum IP66, and luminous efficacy of 120 lm/W. "
            "Relevant standards: IS 10322 (Part 5/Sec 3)."
        ),
        "title": "Municipal Tender",
        "product_description": "LED street light luminaire",
        "sector": "electrical",
    }

    response = client.post("/api/v1/tender/analyze/text", json=payload)
    assert response.status_code == 200

    body = response.json()
    assert len(body["requirements"]) >= 2
    assert body["compliance_matrix"]["product_description"] == "LED street light luminaire"


# ---------------------------------------------------------------------------
# TEST 5: DOCX Report Export
# ---------------------------------------------------------------------------
def test_tender_export_docx_endpoint(client):
    """Verify /api/v1/tender/report/docx returns downloadable Word document."""
    payload = {
        "tender_title": "Municipal Street Lighting RFP",
        "product_description": "90W LED Street Lighting Luminaire",
        "sector": "electrical",
        "compliance_matrix": {
            "overall_grade": "FULLY COMPLIANT",
            "compliance_score": 0.95,
            "evaluations": [
                {
                    "requirement_id": "REQ-001",
                    "parameter": "Ingress Protection (IP)",
                    "specified_value": "IP66",
                    "standard": "IS 10322 (Part 5/Sec 3)",
                    "required_standard_value": "IP65",
                    "status": "PASS",
                }
            ],
            "standards_assessed": [
                {
                    "designation": "IS 10322 (Part 5/Sec 3)",
                    "title": "Luminaires - Particular requirements",
                    "applicability": "APPLICABLE",
                    "is_mandatory_qco": True,
                }
            ],
            "risk_flags": [],
            "recommendations": ["Procure through certified BIS licensees."],
        },
    }

    response = client.post("/api/v1/tender/report/docx", json=payload)
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument")
    assert "attachment; filename=" in response.headers.get("content-disposition", "")
    assert len(response.content) > 1000


# ---------------------------------------------------------------------------
# TEST 6: PDF Report Export
# ---------------------------------------------------------------------------
def test_tender_export_pdf_endpoint(client):
    """Verify /api/v1/tender/report/pdf returns downloadable PDF document."""
    payload = {
        "tender_title": "Municipal Street Lighting RFP",
        "product_description": "90W LED Street Lighting Luminaire",
        "sector": "electrical",
        "compliance_matrix": {
            "overall_grade": "FULLY COMPLIANT",
            "compliance_score": 0.95,
            "evaluations": [
                {
                    "requirement_id": "REQ-001",
                    "parameter": "Ingress Protection (IP)",
                    "specified_value": "IP66",
                    "standard": "IS 10322 (Part 5/Sec 3)",
                    "required_standard_value": "IP65",
                    "status": "PASS",
                }
            ],
            "standards_assessed": [
                {
                    "designation": "IS 10322 (Part 5/Sec 3)",
                    "title": "Luminaires - Particular requirements",
                    "applicability": "APPLICABLE",
                    "is_mandatory_qco": True,
                }
            ],
            "risk_flags": [],
            "recommendations": [],
        },
    }

    response = client.post("/api/v1/tender/report/pdf", json=payload)
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/pdf")
    assert "attachment; filename=" in response.headers.get("content-disposition", "")
    assert response.content.startswith(b"%PDF")
    assert len(response.content) > 1000


# ---------------------------------------------------------------------------
# TEST 7: Empty and Malformed Upload Rejections
# ---------------------------------------------------------------------------
def test_tender_empty_file_rejection(client):
    """Verify empty upload fails safely with HTTP 422."""
    files = {"file": ("empty.pdf", b"", "application/pdf")}
    response = client.post("/api/v1/tender/analyze/file", files=files)
    assert response.status_code == 422


def test_tender_scanned_pdf_rejection(client):
    """Verify image-only / blank PDF fails with HTTP 422 (no OCR)."""
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    c.rect(50, 50, 100, 100)  # drawing without text layer
    c.showPage()
    c.save()

    files = {"file": ("scanned.pdf", buf.getvalue(), "application/pdf")}
    response = client.post("/api/v1/tender/analyze/file", files=files)
    assert response.status_code == 422
    assert "scanned or image-only" in response.json()["detail"] or "usable text layer" in response.json()["detail"]


def test_tender_unsupported_format_rejection(client):
    """Verify unsupported file type is rejected with HTTP 422."""
    files = {"file": ("malicious.exe", b"MZ\x90\x00\x03", "application/octet-stream")}
    response = client.post("/api/v1/tender/analyze/file", files=files)
    assert response.status_code == 422
    assert "Unsupported" in response.json()["detail"]
