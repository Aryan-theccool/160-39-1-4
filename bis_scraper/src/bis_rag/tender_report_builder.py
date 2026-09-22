"""Professional Document Generation for BIS Tender Compliance Assessment Reports (DOCX and PDF).

Generates:
1. DOCX: Native Microsoft Word document with formatted tables and color-coded status chips via python-docx.
2. PDF: Publication-grade A4 PDF report with headers, footers, and compliance matrix via ReportLab.
"""

from __future__ import annotations

import datetime
import io
from typing import Any, Optional

# --- ReportLab Imports ---
import reportlab
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm, inch
from reportlab.pdfgen import canvas
from reportlab.platypus import (
    HRFlowable,
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

# --- docx Imports ---
import docx
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

__all__ = [
    "build_tender_report_docx",
    "build_tender_report_pdf",
    "normalize_tender_report_data",
]


# ======================================================================
# Data Normalization Helper
# ======================================================================
def normalize_tender_report_data(data: Any) -> dict[str, Any]:
    """Convert any input shape into a normalized dictionary for report generation."""
    if hasattr(data, "model_dump"):
        data_dict = data.model_dump()
    elif hasattr(data, "as_dict"):
        data_dict = data.as_dict()
    elif isinstance(data, dict):
        data_dict = data
    else:
        data_dict = {}

    title = str(data_dict.get("tender_title") or "Tender Compliance Assessment Report")
    product_desc = str(data_dict.get("product_description") or "Procurement Item / Specifications")
    sector = data_dict.get("sector") or "General Engineering"

    matrix = data_dict.get("compliance_matrix") or {}
    evaluations = matrix.get("evaluations") or data_dict.get("evaluations") or []
    standards = matrix.get("standards_assessed") or data_dict.get("standards_assessed") or []
    missing_qco = matrix.get("mandatory_qco_missing") or []
    present_qco = matrix.get("mandatory_qco_present") or []
    superseded = matrix.get("superseded_standards_cited") or []
    risk_flags = matrix.get("risk_flags") or []
    recommendations = matrix.get("recommendations") or []
    overall_grade = str(matrix.get("overall_grade") or "UNKNOWN")
    score = float(matrix.get("compliance_score") or 0.0)
    status = str(matrix.get("status") or "UNKNOWN")

    clauses = data_dict.get("clauses") or []
    requirements = data_dict.get("requirements") or []
    detected_standards = data_dict.get("detected_standards") or []

    return {
        "title": title,
        "product_description": product_desc,
        "sector": sector,
        "overall_grade": overall_grade,
        "compliance_score": score,
        "status": status,
        "evaluations": evaluations,
        "standards": standards,
        "missing_qco": missing_qco,
        "present_qco": present_qco,
        "superseded": superseded,
        "risk_flags": risk_flags,
        "recommendations": recommendations,
        "clauses": clauses,
        "requirements": requirements,
        "detected_standards": detected_standards,
        "timestamp": datetime.datetime.now().strftime("%d %B %Y, %I:%M %p"),
        "report_id": f"BIS-TNDR-{datetime.datetime.now().strftime('%Y%m%d-%H%M%S')}",
    }


# ======================================================================
# ReportLab Canvas with Running Header and Footer
# ======================================================================
class NumberedCanvas(canvas.Canvas):
    """Two-pass canvas adding running headers and page count footers."""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count: int):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#718096"))

        # Running Header (pages > 1)
        if self._pageNumber > 1:
            self.drawString(54, A4[1] - 36, "BIS Standards AI — Tender Compliance Assessment Matrix")
            self.setStrokeColor(colors.HexColor("#CBD5E0"))
            self.setLineWidth(0.5)
            self.line(54, A4[1] - 42, A4[0] - 54, A4[1] - 42)

        # Running Footer (all pages)
        self.setStrokeColor(colors.HexColor("#CBD5E0"))
        self.setLineWidth(0.5)
        self.line(54, 45, A4[0] - 54, 45)

        self.drawString(54, 32, "CONFIDENTIAL — FOR PROCUREMENT & COMPLIANCE EVALUATION ONLY")
        page_str = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(A4[0] - 54, 32, page_str)

        self.restoreState()


# ======================================================================
# PDF GENERATION
# ======================================================================
def build_tender_report_pdf(data: Any) -> io.BytesIO:
    """Generate publication-grade PDF tender compliance report using ReportLab."""
    norm = normalize_tender_report_data(data)
    buf = io.BytesIO()

    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54,
    )

    styles = getSampleStyleSheet()

    # Custom Styles
    c_primary = colors.HexColor("#1A365D")
    c_secondary = colors.HexColor("#2B6CB0")
    c_dark = colors.HexColor("#2D3748")

    title_style = ParagraphStyle(
        "TenderTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        textColor=c_primary,
        spaceAfter=4,
    )
    subtitle_style = ParagraphStyle(
        "TenderSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#4A5568"),
        spaceAfter=12,
    )
    h1_style = ParagraphStyle(
        "TenderH1",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=13,
        leading=17,
        textColor=c_primary,
        spaceBefore=12,
        spaceAfter=6,
    )
    body_style = ParagraphStyle(
        "TenderBody",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=13,
        textColor=c_dark,
    )
    bold_body = ParagraphStyle(
        "TenderBoldBody",
        parent=body_style,
        fontName="Helvetica-Bold",
    )

    elements = []

    # Title & Subtitle Header
    elements.append(Paragraph("Tender Compliance & Standards Assessment", title_style))
    elements.append(Paragraph(f"Document: <b>{norm['title']}</b> | Generated: {norm['timestamp']} | Report ID: {norm['report_id']}", subtitle_style))
    elements.append(HRFlowable(width="100%", thickness=1.5, color=c_secondary, spaceAfter=12))

    # Executive Overview Box
    grade_color = "#276749" if "FULLY" in norm["overall_grade"] else ("#D69E2E" if "MOSTLY" in norm["overall_grade"] or "PARTIALLY" in norm["overall_grade"] else "#C53030")
    grade_html = f'<font color="{grade_color}"><b>{norm["overall_grade"]}</b> (Score: {norm["compliance_score"]*100:.1f}%)</font>'

    overview_data = [
        [Paragraph("<b>Target Product:</b>", bold_body), Paragraph(norm["product_description"], body_style)],
        [Paragraph("<b>Procurement Sector:</b>", bold_body), Paragraph(norm["sector"], body_style)],
        [Paragraph("<b>Compliance Verdict:</b>", bold_body), Paragraph(grade_html, body_style)],
        [Paragraph("<b>Total Requirements:</b>", bold_body), Paragraph(str(len(norm["evaluations"])), body_style)],
    ]
    overview_table = Table(overview_data, colWidths=[130, 350])
    overview_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F7FAFC")),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E0")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
    ]))
    elements.append(overview_table)
    elements.append(Spacer(1, 14))

    # Section 1: Clause-by-Clause Compliance Matrix Table
    elements.append(Paragraph("1. Technical Requirements Compliance Matrix", h1_style))
    if norm["evaluations"]:
        matrix_header = [
            Paragraph("<b>Req ID</b>", bold_body),
            Paragraph("<b>Parameter</b>", bold_body),
            Paragraph("<b>Specified Value</b>", bold_body),
            Paragraph("<b>Standard / Required</b>", bold_body),
            Paragraph("<b>Status</b>", bold_body),
        ]
        matrix_rows = [matrix_header]
        for e in norm["evaluations"]:
            st = str(e.get("status", "UNKNOWN"))
            status_color = "#276749" if st == "PASS" else ("#C53030" if st == "FAIL" else "#718096")
            status_chip = f'<font color="{status_color}"><b>{st}</b></font>'

            row = [
                Paragraph(str(e.get("requirement_id", "")), body_style),
                Paragraph(f"<b>{e.get('parameter', '')}</b><br/><font size=7 color='#718096'>{e.get('clause_reference', '')}</font>", body_style),
                Paragraph(str(e.get("specified_value", "")), body_style),
                Paragraph(f"{e.get('standard', '')}<br/><font size=7 color='#4A5568'>{e.get('required_standard_value', '')}</font>", body_style),
                Paragraph(status_chip, body_style),
            ]
            matrix_rows.append(row)

        mat_table = Table(matrix_rows, colWidths=[55, 120, 110, 135, 60])
        mat_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EDF2F7")),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E0")),
            ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ]))
        elements.append(mat_table)
    else:
        elements.append(Paragraph("No technical requirements extracted from the tender document.", body_style))
    elements.append(Spacer(1, 14))

    # Section 2: Standards Assessment & QCO Verification
    elements.append(Paragraph("2. Referenced & Mandatory Standards Assessment", h1_style))
    if norm["standards"]:
        std_header = [
            Paragraph("<b>Designation</b>", bold_body),
            Paragraph("<b>Standard Title</b>", bold_body),
            Paragraph("<b>Applicability</b>", bold_body),
            Paragraph("<b>QCO Mandatory</b>", bold_body),
        ]
        std_rows = [std_header]
        for s in norm["standards"]:
            app = str(s.get("applicability", "CANDIDATE"))
            app_color = "#276749" if app == "APPLICABLE" else ("#C53030" if app == "NOT_APPLICABLE" else "#D69E2E")
            app_html = f'<font color="{app_color}"><b>{app}</b></font>'
            qco_state = s.get("qco_evidence_state", "VERIFIED" if s.get("is_mandatory_qco") else "NOT_FOUND")
            if qco_state == "VERIFIED":
                qco_html = '<font color="#276749"><b>YES (Verified QCO)</b></font>'
            elif qco_state == "SUPERSEDED_OR_CHANGED":
                qco_html = '<font color="#D69E2E"><b>STATUS CHANGED</b></font>'
            elif qco_state == "NOT_APPLICABLE":
                qco_html = '<font color="#718096">Not Applicable</font>'
            elif qco_state == "AMBIGUOUS":
                qco_html = '<font color="#D69E2E">Ambiguous</font>'
            elif qco_state == "NOT_FOUND":
                qco_html = '<font color="#718096">No Gazette QCO Found</font>'
            elif qco_state == "SEED_UNVERIFIED":
                qco_html = '<font color="#D69E2E"><b>YES (Unverified Seed)</b></font>'
            else:
                qco_html = f'<font color="#718096">{qco_state}</font>'

            row = [
                Paragraph(f"<b>{s.get('designation', '')}</b>", body_style),
                Paragraph(str(s.get("title", "")), body_style),
                Paragraph(app_html, body_style),
                Paragraph(qco_html, body_style),
            ]
            std_rows.append(row)

        std_table = Table(std_rows, colWidths=[110, 210, 85, 75])
        std_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EDF2F7")),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E0")),
            ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ]))
        elements.append(std_table)
    elements.append(Spacer(1, 14))

    # Section 3: Risk Flags & Actionable Recommendations
    if norm["risk_flags"] or norm["recommendations"]:
        elements.append(Paragraph("3. Risk Flags & Corrective Recommendations", h1_style))
        for rf in norm["risk_flags"]:
            elements.append(Paragraph(f"<font color='#C53030'><b>[RISK FLAG]</b></font> {rf}", body_style))
            elements.append(Spacer(1, 3))
        for rec in norm["recommendations"]:
            elements.append(Paragraph(f"<font color='#2B6CB0'><b>[ACTION]</b></font> {rec}", body_style))
            elements.append(Spacer(1, 3))

    doc.build(elements, canvasmaker=NumberedCanvas)
    buf.seek(0)
    return buf


# ======================================================================
# DOCX GENERATION
# ======================================================================
def _set_cell_bg(cell, hex_color: str):
    """Set shading color for a DOCX cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def build_tender_report_docx(data: Any) -> io.BytesIO:
    """Generate professional Word (.docx) tender compliance report."""
    norm = normalize_tender_report_data(data)
    doc = docx.Document()

    # Document Geometry
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Title
    p_title = doc.add_paragraph()
    run_title = p_title.add_run("Tender Compliance & Standards Assessment")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    p_title.paragraph_format.space_after = Pt(2)

    # Subtitle
    p_sub = doc.add_paragraph()
    run_sub = p_sub.add_run(f"Document: {norm['title']} | Generated: {norm['timestamp']} | ID: {norm['report_id']}")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(9.5)
    run_sub.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)
    p_sub.paragraph_format.space_after = Pt(12)

    # Overview Table
    tbl_overview = doc.add_table(rows=4, cols=2)
    tbl_overview.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_overview.autofit = False

    overview_items = [
        ("Target Product:", norm["product_description"]),
        ("Procurement Sector:", norm["sector"]),
        ("Overall Grade:", f"{norm['overall_grade']} (Score: {norm['compliance_score']*100:.1f}%)"),
        ("Total Requirements:", str(len(norm["evaluations"]))),
    ]
    for idx, (label, val) in enumerate(overview_items):
        cell_lbl, cell_val = tbl_overview.rows[idx].cells[0], tbl_overview.rows[idx].cells[1]
        cell_lbl.text = label
        cell_lbl.paragraphs[0].runs[0].font.bold = True
        cell_lbl.paragraphs[0].runs[0].font.name = "Arial"
        cell_lbl.paragraphs[0].runs[0].font.size = Pt(9.5)
        _set_cell_bg(cell_lbl, "EDF2F7")

        cell_val.text = val
        cell_val.paragraphs[0].runs[0].font.name = "Arial"
        cell_val.paragraphs[0].runs[0].font.size = Pt(9.5)
        _set_cell_bg(cell_val, "F7FAFC")

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section 1: Compliance Matrix
    h1 = doc.add_heading("1. Technical Requirements Compliance Matrix", level=1)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)

    if norm["evaluations"]:
        headers = ["Req ID", "Parameter", "Specified Value", "Standard / Required", "Status"]
        tbl_matrix = doc.add_table(rows=1 + len(norm["evaluations"]), cols=5)
        tbl_matrix.alignment = WD_TABLE_ALIGNMENT.CENTER

        for col_idx, h in enumerate(headers):
            cell = tbl_matrix.rows[0].cells[col_idx]
            cell.text = h
            cell.paragraphs[0].runs[0].font.bold = True
            cell.paragraphs[0].runs[0].font.name = "Arial"
            cell.paragraphs[0].runs[0].font.size = Pt(9.5)
            _set_cell_bg(cell, "EDF2F7")

        for r_idx, e in enumerate(norm["evaluations"], start=1):
            row = tbl_matrix.rows[r_idx]
            st = str(e.get("status", "UNKNOWN"))

            row.cells[0].text = str(e.get("requirement_id", ""))
            row.cells[1].text = f"{e.get('parameter', '')}\n({e.get('clause_reference', '')})"
            row.cells[2].text = str(e.get("specified_value", ""))
            row.cells[3].text = f"{e.get('standard', '')}\n{e.get('required_standard_value', '')}"
            row.cells[4].text = st

            for c in row.cells:
                if c.paragraphs and c.paragraphs[0].runs:
                    c.paragraphs[0].runs[0].font.name = "Arial"
                    c.paragraphs[0].runs[0].font.size = Pt(9)

            if st == "PASS":
                _set_cell_bg(row.cells[4], "C6F6D5")
            elif st == "FAIL":
                _set_cell_bg(row.cells[4], "FED7D7")
            else:
                _set_cell_bg(row.cells[4], "EDF2F7")

    # Section 2: Standards Assessment
    h2 = doc.add_heading("2. Referenced & Mandatory Standards Assessment", level=1)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)

    if norm["standards"]:
        std_headers = ["Designation", "Standard Title", "Applicability", "QCO Mandatory"]
        tbl_std = doc.add_table(rows=1 + len(norm["standards"]), cols=4)
        tbl_std.alignment = WD_TABLE_ALIGNMENT.CENTER

        for col_idx, h in enumerate(std_headers):
            cell = tbl_std.rows[0].cells[col_idx]
            cell.text = h
            cell.paragraphs[0].runs[0].font.bold = True
            cell.paragraphs[0].runs[0].font.name = "Arial"
            cell.paragraphs[0].runs[0].font.size = Pt(9.5)
            _set_cell_bg(cell, "EDF2F7")

        for r_idx, s in enumerate(norm["standards"], start=1):
            row = tbl_std.rows[r_idx]
            row.cells[0].text = str(s.get("designation", ""))
            row.cells[1].text = str(s.get("title", ""))
            row.cells[2].text = str(s.get("applicability", "CANDIDATE"))
            qco_state = s.get("qco_evidence_state", "VERIFIED" if s.get("is_mandatory_qco") else "NOT_FOUND")
            if qco_state == "VERIFIED":
                row.cells[3].text = "YES (Verified QCO)"
            elif qco_state == "SUPERSEDED_OR_CHANGED":
                row.cells[3].text = "STATUS CHANGED"
            elif qco_state == "NOT_APPLICABLE":
                row.cells[3].text = "Not Applicable"
            elif qco_state == "AMBIGUOUS":
                row.cells[3].text = "Ambiguous"
            elif qco_state == "NOT_FOUND":
                row.cells[3].text = "No Gazette QCO Found"
            elif qco_state == "SEED_UNVERIFIED":
                row.cells[3].text = "YES (Unverified Seed)"
            else:
                row.cells[3].text = str(qco_state)

            for c in row.cells:
                if c.paragraphs and c.paragraphs[0].runs:
                    c.paragraphs[0].runs[0].font.name = "Arial"
                    c.paragraphs[0].runs[0].font.size = Pt(9)

    # Section 3: Risk Flags & Recommendations
    if norm["risk_flags"] or norm["recommendations"]:
        h3 = doc.add_heading("3. Risk Flags & Corrective Recommendations", level=1)
        h3.paragraph_format.space_before = Pt(14)
        h3.paragraph_format.space_after = Pt(6)

        for rf in norm["risk_flags"]:
            p = doc.add_paragraph()
            run_flag = p.add_run("[RISK FLAG] ")
            run_flag.font.bold = True
            run_flag.font.color.rgb = RGBColor(0xC5, 0x30, 0x30)
            p.add_run(rf)

        for rec in norm["recommendations"]:
            p = doc.add_paragraph()
            run_rec = p.add_run("[ACTION] ")
            run_rec.font.bold = True
            run_rec.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
            p.add_run(rec)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
