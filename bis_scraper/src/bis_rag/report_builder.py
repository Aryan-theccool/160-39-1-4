"""Professional Document Generation for BIS Standards Recommendation Reports.

Generates:
1. DOCX: Native Microsoft Word document via python-docx.
2. PDF: Professional A4 PDF via ReportLab (using SimpleDocTemplate, Paragraphs, Tables, and NumberedCanvas).
"""

from __future__ import annotations

import io
import datetime
import re
from typing import Any, Optional

# --- ReportLab Imports ---
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import inch, cm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether, HRFlowable
)
from reportlab.pdfgen import canvas

# --- docx Imports ---
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn


# ======================================================================
# Data Normalization Helper
# ======================================================================
def normalize_report_data(data: Any) -> dict[str, Any]:
    """Convert any input shape (dict, RecommendResponseOut, etc.) into a normalized dict."""
    if hasattr(data, "model_dump"):
        data_dict = data.model_dump()
    elif hasattr(data, "as_dict"):
        data_dict = data.as_dict()
    elif isinstance(data, dict):
        data_dict = data
    else:
        data_dict = {}

    question = str(data_dict.get("question", "") or data_dict.get("query", ""))
    answer = str(data_dict.get("answer", ""))
    report = data_dict.get("report") or {}
    citations = data_dict.get("citations") or []
    warnings = data_dict.get("warnings") or []
    unsupported_citations = data_dict.get("unsupported_citations") or []
    grounded = data_dict.get("grounded", True)
    mode = data_dict.get("mode", "llm")
    generator = data_dict.get("generator", "")

    # Extract report fields
    summary = str(report.get("summary", "") or answer)
    product_und = report.get("product_understanding") or {}
    product_name = str(product_und.get("product", "") or question)
    application = str(product_und.get("application", "Technical Procurement / Tender Evaluation"))
    specs = product_und.get("specifications") or {}

    applicable_standards = report.get("applicable_standards") or []
    related_standards = report.get("related_standards") or []
    compliance_notes = report.get("compliance_notes") or []
    evidence_items = report.get("evidence") or []
    report_warnings = report.get("warnings") or []

    # Merge all warnings without duplicates
    all_warnings = []
    for w in warnings + report_warnings:
        if w and w not in all_warnings:
            all_warnings.append(w)

    return {
        "question": question,
        "product_name": product_name,
        "application": application,
        "specifications": specs,
        "summary": summary,
        "applicable_standards": applicable_standards,
        "related_standards": related_standards,
        "compliance_notes": compliance_notes,
        "warnings": all_warnings,
        "unsupported_citations": unsupported_citations,
        "evidence_items": evidence_items,
        "citations": citations,
        "grounded": grounded,
        "mode": mode,
        "generator": generator,
        "timestamp": datetime.datetime.now().strftime("%d %B %Y, %I:%M %p"),
        "report_id": f"BIS-REC-{datetime.datetime.now().strftime('%Y%m%d-%H%M%S')}",
    }


# ======================================================================
# DOCX Generation Implementation
# ======================================================================
def set_cell_background(cell, fill_hex: str):
    """Set shading color for a DOCX table cell."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set internal cell padding (in twips)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def add_docx_heading(doc: docx.Document, text: str, level: int = 1):
    """Add a styled section heading with government/PSU styling."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    
    run = p.add_run(text)
    run.bold = True
    if level == 1:
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0, 51, 102)  # Navy
        run.font.name = 'Calibri'
    else:
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(51, 65, 85)  # Slate dark
        run.font.name = 'Calibri'
    return p


def build_report_docx(data_input: Any) -> io.BytesIO:
    """Build a complete, professionally formatted Microsoft Word (.docx) document."""
    data = normalize_report_data(data_input)
    doc = docx.Document()

    # Set page margins to standard A4 (1 inch = 72 pt)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        
        # Header / Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Bureau of Indian Standards — Procurement Recommendation Report")
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(128, 128, 128)

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        frun = fp.add_run(f"Report ID: {data['report_id']} | Generated: {data['timestamp']}")
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(128, 128, 128)

    # 1. Main Document Header
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run("AI-Powered Indian Standards Recommendation Report")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(18)
    run_title.bold = True
    run_title.font.color.rgb = RGBColor(0, 51, 102)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(10)
    run_sub = p_sub.add_run("Procurement Standards Assessment & Technical Compliance Screening")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = RGBColor(100, 116, 139)

    # Metadata Box (Table)
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_rows = [
        ("Target Query / Product:", data["product_name"] or data["question"]),
        ("Report Identifier:", data["report_id"]),
        ("Date & Time of Generation:", data["timestamp"]),
        ("Assessment Scope & Grounding:", 
         "Grounded on Bureau of Indian Standards indexed partition dataset. "
         + ("All citations verified." if data["grounded"] else "Contains unverified external citations requiring review.")),
    ]

    col_widths = [Inches(2.2), Inches(4.4)]
    for r_idx, (label, val) in enumerate(meta_rows):
        row = meta_table.rows[r_idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width, cell_val.width = col_widths[0], col_widths[1]

        p_l = cell_lbl.paragraphs[0]
        p_l.paragraph_format.space_before = Pt(2)
        p_l.paragraph_format.space_after = Pt(2)
        r_l = p_l.add_run(label)
        r_l.bold = True
        r_l.font.size = Pt(9.5)
        r_l.font.color.rgb = RGBColor(0, 51, 102)

        p_v = cell_val.paragraphs[0]
        p_v.paragraph_format.space_before = Pt(2)
        p_v.paragraph_format.space_after = Pt(2)
        r_v = p_v.add_run(val)
        r_v.font.size = Pt(9.5)
        r_v.font.color.rgb = RGBColor(30, 41, 59)

        set_cell_background(cell_lbl, "F8FAFC")
        set_cell_background(cell_val, "FFFFFF")
        set_cell_margins(cell_lbl, top=80, bottom=80, left=120, right=120)
        set_cell_margins(cell_val, top=80, bottom=80, left=120, right=120)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 2. Section 1: Executive Summary
    add_docx_heading(doc, "1. Executive Summary", level=1)
    p_sum = doc.add_paragraph()
    p_sum.paragraph_format.space_after = Pt(8)
    p_sum.paragraph_format.line_spacing = 1.15
    r_sum = p_sum.add_run(data["summary"] or "No summary available.")
    r_sum.font.size = Pt(10)
    r_sum.font.name = 'Calibri'

    # 3. Section 2: Product & Specification Understanding
    add_docx_heading(doc, "2. Product & Procurement Understanding", level=1)
    p_pu = doc.add_paragraph()
    p_pu.paragraph_format.space_after = Pt(4)
    r_pu1 = p_pu.add_run("• Intended Product / Scope: ")
    r_pu1.bold = True
    r_pu1.font.size = Pt(10)
    p_pu.add_run(data["product_name"]).font.size = Pt(10)

    p_app = doc.add_paragraph()
    p_app.paragraph_format.space_after = Pt(6)
    r_app1 = p_app.add_run("• Target Application / Domain: ")
    r_app1.bold = True
    r_app1.font.size = Pt(10)
    p_app.add_run(data["application"]).font.size = Pt(10)

    if data["specifications"]:
        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_after = Pt(4)
        r_sp = p_sp.add_run("• Extracted Parameters & Specifications:")
        r_sp.bold = True
        r_sp.font.size = Pt(10)
        for k, v in data["specifications"].items():
            p_kv = doc.add_paragraph()
            p_kv.paragraph_format.left_indent = Inches(0.25)
            p_kv.paragraph_format.space_after = Pt(2)
            p_kv.add_run(f"- {k}: ").bold = True
            p_kv.add_run(str(v))
            for r in p_kv.runs:
                r.font.size = Pt(9.5)

    # 4. Section 3: Applicable Indian Standards Table
    add_docx_heading(doc, "3. Applicable Indian Standards", level=1)
    if data["applicable_standards"]:
        app_table = doc.add_table(rows=len(data["applicable_standards"]) + 1, cols=6)
        app_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        app_table.autofit = False

        headers = ["#", "Standard", "Title", "Status", "Relevance", "Applicability / Reason"]
        widths = [Inches(0.4), Inches(1.2), Inches(1.8), Inches(0.8), Inches(0.9), Inches(1.5)]
        
        # Header row
        hdr_row = app_table.rows[0]
        for c_idx, h_text in enumerate(headers):
            cell = hdr_row.cells[c_idx]
            cell.width = widths[c_idx]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in (0, 3, 4) else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(h_text)
            r.bold = True
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(255, 255, 255)
            set_cell_background(cell, "003366")
            set_cell_margins(cell, top=100, bottom=100, left=100, right=100)

        # Data rows
        for idx, item in enumerate(data["applicable_standards"], start=1):
            row = app_table.rows[idx]
            desig = str(item.get("designation", ""))
            title = str(item.get("title", ""))
            status = str(item.get("status", "current")).capitalize()
            rel = str(item.get("relevance", "Applicable"))
            reason = str(item.get("reason", ""))

            vals = [str(idx), desig, title, status, rel, reason]
            for c_idx, val in enumerate(vals):
                cell = row.cells[c_idx]
                cell.width = widths[c_idx]
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in (0, 3, 4) else WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(val)
                r.font.size = Pt(8.5)
                if c_idx == 1:
                    r.bold = True
                if c_idx == 3:
                    r.font.color.rgb = RGBColor(16, 185, 129) if status.lower() == "current" else RGBColor(217, 119, 6)

                bg = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
                set_cell_background(cell, bg)
                set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
    else:
        p_none = doc.add_paragraph("No specific applicable standards retrieved in this dataset partition.")
        p_none.runs[0].font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 5. Section 4: Related & Allied Standards
    add_docx_heading(doc, "4. Related & Allied Standards", level=1)
    if data["related_standards"]:
        rel_table = doc.add_table(rows=len(data["related_standards"]) + 1, cols=4)
        rel_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        rel_table.autofit = False

        r_headers = ["#", "Standard", "Title", "Relationship / Scope"]
        r_widths = [Inches(0.4), Inches(1.4), Inches(2.8), Inches(2.0)]
        for c_idx, h_text in enumerate(r_headers):
            cell = rel_table.rows[0].cells[c_idx]
            cell.width = r_widths[c_idx]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(h_text)
            r.bold = True
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(255, 255, 255)
            set_cell_background(cell, "334155")
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)

        for idx, item in enumerate(data["related_standards"], start=1):
            row = rel_table.rows[idx]
            vals = [str(idx), str(item.get("designation", "")), str(item.get("title", "")), str(item.get("relationship", "Related"))]
            for c_idx, val in enumerate(vals):
                cell = row.cells[c_idx]
                cell.width = r_widths[c_idx]
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(val)
                r.font.size = Pt(8.5)
                if c_idx == 1:
                    r.bold = True
                set_cell_background(cell, "F8FAFC" if idx % 2 == 1 else "FFFFFF")
                set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
    else:
        p_none = doc.add_paragraph("No allied or supporting standards identified.")
        p_none.runs[0].font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 6. Section 5: Compliance & Quality Notes
    add_docx_heading(doc, "5. Compliance & Quality Notes", level=1)
    if data["compliance_notes"]:
        for note in data["compliance_notes"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(f"• {note}")
            r.font.size = Pt(9.5)
    else:
        p = doc.add_paragraph("• Quality Control Order (QCO) and mandatory certification status should be verified against the official Gazette of India notifications.")
        p.runs[0].font.size = Pt(9.5)

    # 7. Section 6: Warnings & Grounding Alerts
    add_docx_heading(doc, "6. Warnings & Grounding Alerts", level=1)
    if data["warnings"] or data["unsupported_citations"]:
        w_box = doc.add_table(rows=1, cols=1)
        w_box.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = w_box.rows[0].cells[0]
        cell.width = Inches(6.6)
        set_cell_background(cell, "FEF3C7")  # Amber light
        set_cell_margins(cell, top=100, bottom=100, left=150, right=150)

        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(3)
        r_head = p.add_run("ATTENTION / VERIFICATION REQUIRED:\n")
        r_head.bold = True
        r_head.font.size = Pt(9.5)
        r_head.font.color.rgb = RGBColor(180, 83, 9)

        for w in data["warnings"]:
            p_w = cell.add_paragraph()
            p_w.paragraph_format.space_after = Pt(2)
            r = p_w.add_run(f"• {w}")
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(146, 64, 14)

        if data["unsupported_citations"]:
            for un_desig in data["unsupported_citations"]:
                p_u = cell.add_paragraph()
                p_u.paragraph_format.space_after = Pt(2)
                r_u = p_u.add_run(f"• Unverified standard reference: {un_desig}. This standard was mentioned by the model but was not available in the indexed knowledge base.")
                r_u.font.size = Pt(9)
                r_u.font.color.rgb = RGBColor(146, 64, 14)
    else:
        p_ok = doc.add_paragraph("✓ All cited Indian Standards verified against the indexed evidence corpus with no supersession conflicts.")
        p_ok.runs[0].font.size = Pt(9.5)
        p_ok.runs[0].font.color.rgb = RGBColor(16, 185, 129)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 8. Section 7: Evidence & Citations
    add_docx_heading(doc, "7. Evidence & Public Standard References", level=1)
    if data["citations"]:
        for c in data["citations"]:
            desig = c.get("designation", "")
            title = c.get("title", "")
            url = c.get("archive_url", "")
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            r_d = p.add_run(f"• {desig}: ")
            r_d.bold = True
            r_d.font.size = Pt(9)
            p.add_run(f"{title} ").font.size = Pt(9)
            if url:
                r_url = p.add_run(f"[{url}]")
                r_url.font.size = Pt(8.5)
                r_url.font.color.rgb = RGBColor(0, 102, 204)
    else:
        p_ev = doc.add_paragraph("No explicit citation links recorded.")
        p_ev.runs[0].font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 9. Section 8: Important Limitations & Disclaimer
    add_docx_heading(doc, "8. Important Limitations & Regulatory Notice", level=1)
    disclaimers = [
        "Corpus Scope: The recommendations in this report are based on the indexed subset of Bureau of Indian Standards (197 partition standards / 9,031 chunked passages). Standards outside this partition may also apply.",
        "QCO & Mandatory Status: Quality Control Orders are legal instruments issued by relevant Ministries. Users must cross-verify mandatory certification requirements with the latest official BIS Gazette notifications.",
        "Regulatory Disclaimer: This document is an AI-assisted technical screening aid for procurement specifications and does not constitute a formal legal opinion or official Bureau of Indian Standards certification.",
    ]
    for disc in disclaimers:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f"• {disc}")
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(100, 116, 139)

    # Save to BytesIO
    out = io.BytesIO()
    doc.save(out)
    out.seek(0)
    return out


# ======================================================================
# PDF Generation Implementation via ReportLab
# ======================================================================
class NumberedCanvas(canvas.Canvas):
    """Two-pass canvas to calculate total page count and draw running header/footer."""
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, total_pages: int):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748B"))

        # Header (pages > 1)
        if self._pageNumber > 1:
            self.drawString(54, 800, "Bureau of Indian Standards — Procurement Standards Recommendation Report")
            self.setStrokeColor(colors.HexColor("#CBD5E1"))
            self.setLineWidth(0.5)
            self.line(54, 792, 541, 792)

        # Footer (all pages)
        self.setStrokeColor(colors.HexColor("#CBD5E1"))
        self.setLineWidth(0.5)
        self.line(54, 45, 541, 45)
        
        self.drawString(54, 32, "Confidential — AI-Powered Technical Procurement Screening")
        page_str = f"Page {self._pageNumber} of {total_pages}"
        self.drawRightString(541, 32, page_str)

        self.restoreState()


def build_report_pdf(data_input: Any) -> io.BytesIO:
    """Build a complete, professionally formatted PDF document via ReportLab."""
    data = normalize_report_data(data_input)
    buffer = io.BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54,
    )

    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        'RepTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=16,
        leading=20,
        textColor=colors.HexColor("#003366"),
        spaceAfter=2,
    )
    subtitle_style = ParagraphStyle(
        'RepSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#64748B"),
        spaceAfter=10,
    )
    h1_style = ParagraphStyle(
        'RepH1',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11.5,
        leading=15,
        textColor=colors.HexColor("#003366"),
        spaceBefore=12,
        spaceAfter=4,
        keepWithNext=True,
    )
    body_style = ParagraphStyle(
        'RepBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12.5,
        textColor=colors.HexColor("#1E293B"),
        spaceAfter=4,
    )
    bullet_style = ParagraphStyle(
        'RepBullet',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11.5,
        textColor=colors.HexColor("#1E293B"),
        leftIndent=12,
        firstLineIndent=-8,
        spaceAfter=3,
    )
    table_cell_style = ParagraphStyle(
        'RepTableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=10.5,
        textColor=colors.HexColor("#1E293B"),
    )
    table_hdr_style = ParagraphStyle(
        'RepTableHdr',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10.5,
        textColor=colors.white,
    )
    warn_style = ParagraphStyle(
        'RepWarn',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11.5,
        textColor=colors.HexColor("#92400E"),
        spaceAfter=2,
    )

    story = []

    # Title & Subtitle
    story.append(Paragraph("AI-Powered Indian Standards Recommendation Report", title_style))
    story.append(Paragraph("Procurement Standards Assessment & Technical Compliance Screening", subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#003366"), spaceAfter=8))

    # Metadata Table
    meta_data = [
        [Paragraph("<b>Target Query / Product:</b>", table_cell_style), Paragraph(data["product_name"] or data["question"], table_cell_style)],
        [Paragraph("<b>Report Identifier:</b>", table_cell_style), Paragraph(data["report_id"], table_cell_style)],
        [Paragraph("<b>Date & Time:</b>", table_cell_style), Paragraph(data["timestamp"], table_cell_style)],
        [Paragraph("<b>Scope & Grounding:</b>", table_cell_style), 
         Paragraph("Grounded on Bureau of Indian Standards dataset partition. " + ("All citations verified." if data["grounded"] else "Contains unverified external citations requiring review."), table_cell_style)],
    ]
    meta_table = Table(meta_data, colWidths=[130, 357])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor("#F8FAFC")),
        ('BACKGROUND', (1, 0), (1, -1), colors.white),
        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 8))

    # 1. Executive Summary
    story.append(Paragraph("1. Executive Summary", h1_style))
    story.append(Paragraph(data["summary"] or "No summary available.", body_style))
    story.append(Spacer(1, 4))

    # 2. Product & Procurement Understanding
    story.append(Paragraph("2. Product & Procurement Understanding", h1_style))
    story.append(Paragraph(f"<b>• Intended Product:</b> {data['product_name']}", bullet_style))
    story.append(Paragraph(f"<b>• Target Domain:</b> {data['application']}", bullet_style))
    if data["specifications"]:
        story.append(Paragraph("<b>• Extracted Parameters:</b>", bullet_style))
        for k, v in data["specifications"].items():
            story.append(Paragraph(f"&nbsp;&nbsp;&nbsp;&nbsp;<b>{k}:</b> {v}", bullet_style))
    story.append(Spacer(1, 4))

    # 3. Applicable Indian Standards
    story.append(Paragraph("3. Applicable Indian Standards", h1_style))
    if data["applicable_standards"]:
        app_rows = [[
            Paragraph("#", table_hdr_style),
            Paragraph("Standard", table_hdr_style),
            Paragraph("Title", table_hdr_style),
            Paragraph("Status", table_hdr_style),
            Paragraph("Relevance", table_hdr_style),
            Paragraph("Applicability / Reason", table_hdr_style),
        ]]
        for idx, item in enumerate(data["applicable_standards"], start=1):
            desig = item.get("designation", "")
            title = item.get("title", "")
            status = str(item.get("status", "current")).capitalize()
            rel = item.get("relevance", "Applicable")
            reason = item.get("reason", "")
            app_rows.append([
                Paragraph(str(idx), table_cell_style),
                Paragraph(f"<b>{desig}</b>", table_cell_style),
                Paragraph(title, table_cell_style),
                Paragraph(f"<font color='{'#10B981' if status.lower()=='current' else '#D97706'}'><b>{status}</b></font>", table_cell_style),
                Paragraph(rel, table_cell_style),
                Paragraph(reason, table_cell_style),
            ])
        app_table = Table(app_rows, colWidths=[20, 85, 140, 50, 62, 130], repeatRows=1)
        app_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#003366")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('ALIGN', (0, 0), (0, -1), 'CENTER'),
            ('ALIGN', (3, 0), (4, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor("#F8FAFC"), colors.white]),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('LEFTPADDING', (0, 0), (-1, -1), 4),
            ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(app_table)
    else:
        story.append(Paragraph("No specific applicable standards retrieved in this dataset partition.", body_style))
    story.append(Spacer(1, 6))

    # 4. Related & Allied Standards
    story.append(Paragraph("4. Related & Allied Standards", h1_style))
    if data["related_standards"]:
        rel_rows = [[
            Paragraph("#", table_hdr_style),
            Paragraph("Standard", table_hdr_style),
            Paragraph("Title", table_hdr_style),
            Paragraph("Relationship / Scope", table_hdr_style),
        ]]
        for idx, item in enumerate(data["related_standards"], start=1):
            rel_rows.append([
                Paragraph(str(idx), table_cell_style),
                Paragraph(f"<b>{item.get('designation', '')}</b>", table_cell_style),
                Paragraph(item.get("title", ""), table_cell_style),
                Paragraph(item.get("relationship", "Related"), table_cell_style),
            ])
        rel_table = Table(rel_rows, colWidths=[20, 100, 227, 140], repeatRows=1)
        rel_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#334155")),
            ('ALIGN', (0, 0), (0, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor("#F8FAFC"), colors.white]),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
            ('LEFTPADDING', (0, 0), (-1, -1), 4),
            ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(rel_table)
    else:
        story.append(Paragraph("No allied or supporting standards identified.", body_style))
    story.append(Spacer(1, 6))

    # 5. Compliance & Quality Notes
    story.append(Paragraph("5. Compliance & Quality Notes", h1_style))
    if data["compliance_notes"]:
        for note in data["compliance_notes"]:
            story.append(Paragraph(f"• {note}", bullet_style))
    else:
        story.append(Paragraph("• Quality Control Order (QCO) and mandatory certification status should be verified against the official Gazette of India notifications.", bullet_style))
    story.append(Spacer(1, 4))

    # 6. Warnings & Grounding Alerts
    story.append(Paragraph("6. Warnings & Grounding Alerts", h1_style))
    if data["warnings"] or data["unsupported_citations"]:
        warn_content = [
            Paragraph("<b>ATTENTION / VERIFICATION REQUIRED:</b>", warn_style)
        ]
        for w in data["warnings"]:
            warn_content.append(Paragraph(f"• {w}", warn_style))
        if data["unsupported_citations"]:
            for un_desig in data["unsupported_citations"]:
                warn_content.append(Paragraph(f"• Unverified standard reference: {un_desig}. This standard was mentioned by the model but was not available in the indexed knowledge base.", warn_style))
        
        warn_table = Table([[warn_content]], colWidths=[487])
        warn_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#FEF3C7")),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor("#F59E0B")),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('LEFTPADDING', (0, 0), (-1, -1), 8),
            ('RIGHTPADDING', (0, 0), (-1, -1), 8),
        ]))
        story.append(warn_table)
    else:
        story.append(Paragraph("✓ <font color='#10B981'><b>All citations verified against indexed evidence corpus with no supersession conflicts.</b></font>", body_style))
    story.append(Spacer(1, 6))

    # 7. Evidence & References
    story.append(Paragraph("7. Evidence & Public Standard References", h1_style))
    if data["citations"]:
        for c in data["citations"]:
            desig = c.get("designation", "")
            title = c.get("title", "")
            url = c.get("archive_url", "")
            url_text = f"<br/><font color='#0066CC'><u>{url}</u></font>" if url else ""
            story.append(Paragraph(f"• <b>{desig}:</b> {title} {url_text}", bullet_style))
    else:
        story.append(Paragraph("No explicit citation links recorded.", body_style))
    story.append(Spacer(1, 6))

    # 8. Limitations & Regulatory Notice
    story.append(Paragraph("8. Limitations & Regulatory Notice", h1_style))
    disclaimers = [
        "<b>Corpus Scope:</b> The recommendations in this report are based on the indexed partition of Bureau of Indian Standards (197 partition standards / 9,031 chunked passages). Standards outside this partition may also apply.",
        "<b>QCO & Mandatory Status:</b> Quality Control Orders are legal instruments issued by relevant Ministries. Users must cross-verify mandatory certification requirements with the latest official BIS Gazette notifications.",
        "<b>Regulatory Disclaimer:</b> This document is an AI-assisted technical screening aid for procurement specifications and does not constitute a formal legal opinion or official Bureau of Indian Standards certification.",
    ]
    for disc in disclaimers:
        story.append(Paragraph(f"• {disc}", ParagraphStyle('RepDisc', parent=styles['Normal'], fontName='Helvetica', fontSize=8, leading=10.5, textColor=colors.HexColor("#64748B"), leftIndent=10, firstLineIndent=-6, spaceAfter=2)))

    # Build PDF with NumberedCanvas
    doc.build(story, canvasmaker=NumberedCanvas)
    buffer.seek(0)
    return buffer
