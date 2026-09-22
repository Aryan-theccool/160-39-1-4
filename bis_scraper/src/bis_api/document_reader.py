"""Unified document reader supporting PDF, DOCX, and TXT formats for tender ingestion.

Preserves structural document metadata (pages, paragraphs with styles, table grids, line boundaries)
needed by downstream Tender Analysis phases while enforcing strict validation, size limits,
and safe handling of scanned/image-only PDFs without OCR.
"""

from __future__ import annotations

import io
import logging
import os
import re
from dataclasses import dataclass, field
from typing import Optional

from .pdf_text import MIN_USABLE_CHARS, _extract_builtin

log = logging.getLogger(__name__)

__all__ = [
    "DEFAULT_MAX_BYTES",
    "DEFAULT_MAX_CHARS",
    "DocumentSource",
    "DocumentPage",
    "DocumentParagraph",
    "DocumentTable",
    "DocumentContent",
    "DocumentReadResult",
    "read_document",
    "read_pdf",
    "read_docx",
    "read_txt",
    "detect_format",
]

DEFAULT_MAX_BYTES = 10 * 1024 * 1024  # 10 MB
DEFAULT_MAX_CHARS = 500_000


@dataclass
class DocumentSource:
    filename: str = ""
    media_type: str = ""
    extension: str = ""
    source_type: str = ""  # "pdf", "docx", "txt", "unknown"

    def as_dict(self) -> dict:
        return {
            "filename": self.filename,
            "media_type": self.media_type,
            "extension": self.extension,
            "source_type": self.source_type,
        }


@dataclass
class DocumentPage:
    page_number: int  # 1-indexed
    text: str
    character_count: int

    def as_dict(self) -> dict:
        return {
            "page_number": self.page_number,
            "text": self.text,
            "character_count": self.character_count,
        }


@dataclass
class DocumentParagraph:
    text: str
    style: str = "Normal"
    paragraph_index: int = 0

    def as_dict(self) -> dict:
        return {
            "text": self.text,
            "style": self.style,
            "paragraph_index": self.paragraph_index,
        }


@dataclass
class DocumentTable:
    table_index: int
    rows: list[list[str]] = field(default_factory=list)
    source_page: Optional[int] = None

    def as_dict(self) -> dict:
        return {
            "table_index": self.table_index,
            "rows": self.rows,
            "source_page": self.source_page,
        }


@dataclass
class DocumentContent:
    source: DocumentSource
    pages: list[DocumentPage] = field(default_factory=list)
    paragraphs: list[DocumentParagraph] = field(default_factory=list)
    tables: list[DocumentTable] = field(default_factory=list)
    full_text: str = ""
    character_count: int = 0
    page_count: int = 0
    extraction_method: str = ""
    warnings: list[str] = field(default_factory=list)

    def as_dict(self) -> dict:
        return {
            "source": self.source.as_dict(),
            "pages": [p.as_dict() for p in self.pages],
            "paragraphs": [p.as_dict() for p in self.paragraphs],
            "tables": [t.as_dict() for t in self.tables],
            "full_text": self.full_text,
            "character_count": self.character_count,
            "page_count": self.page_count,
            "extraction_method": self.extraction_method,
            "warnings": self.warnings,
        }


@dataclass
class DocumentReadResult:
    ok: bool
    document: Optional[DocumentContent] = None
    error_code: str = ""
    error_message: str = ""
    warnings: list[str] = field(default_factory=list)

    def as_dict(self) -> dict:
        return {
            "ok": self.ok,
            "document": self.document.as_dict() if self.document else None,
            "error_code": self.error_code,
            "error_message": self.error_message,
            "warnings": self.warnings,
        }


def detect_format(data: bytes, filename: str = "", media_type: str = "") -> str:
    """Determine document format ('pdf', 'docx', 'txt', or 'unknown') from hints and content."""
    # 1. Check filename extension
    if filename:
        _, ext = os.path.splitext(filename.lower())
        if ext == ".pdf":
            return "pdf"
        if ext in (".docx", ".doc"):  # docx is primary
            return "docx"
        if ext in (".txt", ".text", ".md", ".csv", ".tsv"):
            return "txt"

    # 2. Check media type
    mt = (media_type or "").lower().split(";")[0].strip()
    if mt in ("application/pdf", "application/x-pdf"):
        return "pdf"
    if mt in (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/msword",
        "application/docx",
    ):
        return "docx"
    if mt in ("text/plain", "text/markdown", "text/csv"):
        return "txt"

    # 3. Magic byte inspection
    if data:
        trimmed = data.lstrip()
        if trimmed.startswith(b"%PDF"):
            return "pdf"
        if data.startswith(b"PK\x03\x04"):
            return "docx"
        # If it looks like printable ASCII/UTF-8 text, treat as txt
        try:
            sample = data[:4096].decode("utf-8")
            if sample.isprintable() or any(c in "\r\n\t" for c in sample):
                return "txt"
        except UnicodeDecodeError:
            pass

    return "unknown"


def _extract_pdf_pymupdf(data: bytes, max_chars: int) -> tuple[bool, list[DocumentPage], int, str, list[str]]:
    """Attempt extraction using PyMuPDF."""
    try:
        import pymupdf as fitz
    except ImportError:
        try:
            import fitz  # type: ignore[no-redef]
        except ImportError:
            return False, [], 0, "", []

    warnings: list[str] = []
    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception:
        return False, [], 0, "", []

    try:
        if doc.needs_pass:
            return False, [], doc.page_count, "PASSWORD_PROTECTED", ["PDF is password-protected."]
        pages: list[DocumentPage] = []
        page_count = doc.page_count
        total_chars = 0
        for idx, page in enumerate(doc):
            text = page.get_text() or ""
            text = text.replace("\r\n", "\n").replace("\r", "\n")
            if total_chars + len(text) > max_chars:
                text = text[: max(0, max_chars - total_chars)]
                warnings.append(f"Text truncated at maximum limit of {max_chars} characters.")
            p_chars = len(text)
            total_chars += p_chars
            pages.append(DocumentPage(page_number=idx + 1, text=text, character_count=p_chars))
            if total_chars >= max_chars:
                break
        return True, pages, page_count, "pymupdf", warnings
    finally:
        doc.close()


def _extract_pdf_pypdf(data: bytes, max_chars: int) -> tuple[bool, list[DocumentPage], int, str, list[str]]:
    """Attempt extraction using pypdf."""
    try:
        from pypdf import PdfReader
    except ImportError:
        return False, [], 0, "", []

    warnings: list[str] = []
    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception:
        return False, [], 0, "", []

    if reader.is_encrypted:
        try:
            # Try empty password
            reader.decrypt("")
        except Exception:
            return False, [], len(reader.pages), "PASSWORD_PROTECTED", ["PDF is password-protected."]

    pages: list[DocumentPage] = []
    page_count = len(reader.pages)
    total_chars = 0
    for idx, page in enumerate(reader.pages):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        text = text.replace("\r\n", "\n").replace("\r", "\n")
        if total_chars + len(text) > max_chars:
            text = text[: max(0, max_chars - total_chars)]
            warnings.append(f"Text truncated at maximum limit of {max_chars} characters.")
        p_chars = len(text)
        total_chars += p_chars
        pages.append(DocumentPage(page_number=idx + 1, text=text, character_count=p_chars))
        if total_chars >= max_chars:
            break

    return True, pages, page_count, "pypdf", warnings


def read_pdf(
    data: bytes,
    filename: str = "",
    media_type: str = "",
    *,
    max_chars: int = DEFAULT_MAX_CHARS,
) -> DocumentReadResult:
    """Read a PDF document with per-page tracking, rejecting image-only/scanned PDFs without OCR."""
    source = DocumentSource(
        filename=filename,
        media_type=media_type or "application/pdf",
        extension=".pdf",
        source_type="pdf",
    )
    warnings: list[str] = []

    if not data:
        return DocumentReadResult(
            ok=False,
            error_code="EMPTY_FILE",
            error_message="The uploaded PDF file is empty.",
            warnings=warnings,
        )

    if not data.lstrip()[:5].startswith(b"%PDF"):
        return DocumentReadResult(
            ok=False,
            error_code="MALFORMED_DOCUMENT",
            error_message="The uploaded file is not a valid PDF (missing %PDF header).",
            warnings=warnings,
        )

    pages: list[DocumentPage] = []
    page_count = 0
    method = "none"

    # Try PyMuPDF first
    ok, p_pages, p_count, p_method, p_warn = _extract_pdf_pymupdf(data, max_chars)
    if ok:
        pages, page_count, method = p_pages, p_count, p_method
        warnings.extend(p_warn)
    elif p_method == "PASSWORD_PROTECTED":
        return DocumentReadResult(
            ok=False,
            error_code="PASSWORD_PROTECTED",
            error_message="The PDF is password-protected. Please remove the password and retry.",
            warnings=warnings,
        )
    else:
        # Try pypdf next
        ok2, p2_pages, p2_count, p2_method, p2_warn = _extract_pdf_pypdf(data, max_chars)
        if ok2:
            pages, page_count, method = p2_pages, p2_count, p2_method
            warnings.extend(p2_warn)
        elif p2_method == "PASSWORD_PROTECTED":
            return DocumentReadResult(
                ok=False,
                error_code="PASSWORD_PROTECTED",
                error_message="The PDF is password-protected. Please remove the password and retry.",
                warnings=warnings,
            )

    # If no pages or extracted text is below minimum usable threshold, try builtin stream fallback
    if not pages or sum(p.character_count for p in pages) < MIN_USABLE_CHARS:
        builtin_res = _extract_builtin(data, max_chars=max_chars)
        if builtin_res.ok and len(builtin_res.text) >= MIN_USABLE_CHARS:
            warnings.extend(builtin_res.warnings)
            warnings.append("Extracted using built-in PDF stream reader (no per-page boundaries).")
            full_text = builtin_res.text
            char_count = len(full_text)
            pages = [DocumentPage(page_number=1, text=full_text, character_count=char_count)]
            page_count = max(page_count, 1)
            method = "builtin"
        else:
            return DocumentReadResult(
                ok=False,
                error_code="NO_TEXT_LAYER",
                error_message=(
                    "No usable text layer found in PDF (document appears to be scanned or image-only). "
                    "OCR is not attempted. Please upload a searchable, text-based PDF."
                ),
                warnings=warnings,
            )

    full_text = "\n\n".join(p.text.strip() for p in pages if p.text.strip()).strip()
    char_count = len(full_text)

    if char_count < MIN_USABLE_CHARS:
        return DocumentReadResult(
            ok=False,
            error_code="NO_TEXT_LAYER",
            error_message="No usable text layer found in PDF (insufficient character count).",
            warnings=warnings,
        )

    # Generate logical paragraphs from pages
    paragraphs: list[DocumentParagraph] = []
    p_idx = 0
    for page in pages:
        for block in page.text.split("\n\n"):
            cleaned = block.strip()
            if cleaned:
                paragraphs.append(DocumentParagraph(text=cleaned, style="Normal", paragraph_index=p_idx))
                p_idx += 1

    content = DocumentContent(
        source=source,
        pages=pages,
        paragraphs=paragraphs,
        tables=[],  # Phase 3A preserves empty table list for standard PDF text
        full_text=full_text,
        character_count=char_count,
        page_count=page_count,
        extraction_method=method,
        warnings=warnings,
    )
    return DocumentReadResult(ok=True, document=content, warnings=warnings)


def read_docx(
    data: bytes,
    filename: str = "",
    media_type: str = "",
    *,
    max_chars: int = DEFAULT_MAX_CHARS,
) -> DocumentReadResult:
    """Read a DOCX document preserving paragraph styles and table row/cell structure."""
    source = DocumentSource(
        filename=filename,
        media_type=media_type or "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        extension=".docx",
        source_type="docx",
    )
    warnings: list[str] = []

    if not data:
        return DocumentReadResult(
            ok=False,
            error_code="EMPTY_FILE",
            error_message="The uploaded DOCX file is empty.",
            warnings=warnings,
        )

    try:
        import docx
    except ImportError:
        return DocumentReadResult(
            ok=False,
            error_code="DEPENDENCY_MISSING",
            error_message="python-docx library is not installed.",
            warnings=warnings,
        )

    try:
        doc = docx.Document(io.BytesIO(data))
    except Exception as exc:
        return DocumentReadResult(
            ok=False,
            error_code="MALFORMED_DOCUMENT",
            error_message=f"Failed to parse DOCX document: {exc}",
            warnings=warnings,
        )

    paragraphs: list[DocumentParagraph] = []
    for idx, p in enumerate(doc.paragraphs):
        p_text = p.text.strip()
        style_name = p.style.name if p.style and hasattr(p.style, "name") else "Normal"
        if p_text:
            paragraphs.append(
                DocumentParagraph(
                    text=p_text,
                    style=style_name or "Normal",
                    paragraph_index=idx,
                )
            )

    tables: list[DocumentTable] = []
    for t_idx, table in enumerate(doc.tables):
        rows: list[list[str]] = []
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells]
            rows.append(row_cells)
        if rows:
            tables.append(DocumentTable(table_index=t_idx, rows=rows))

    # Construct unified full_text
    text_chunks: list[str] = []
    for p in paragraphs:
        text_chunks.append(p.text)

    # Append structured table text representation
    for t in tables:
        table_lines = ["\t".join(row) for row in t.rows if any(cell for cell in row)]
        if table_lines:
            text_chunks.append("\n".join(table_lines))

    full_text = "\n\n".join(text_chunks).strip()
    if len(full_text) > max_chars:
        full_text = full_text[:max_chars]
        warnings.append(f"Text truncated at maximum limit of {max_chars} characters.")

    char_count = len(full_text)

    if char_count < MIN_USABLE_CHARS and not tables:
        return DocumentReadResult(
            ok=False,
            error_code="EMPTY_DOCUMENT",
            error_message="The DOCX document contains no readable text or tables.",
            warnings=warnings,
        )

    content = DocumentContent(
        source=source,
        pages=[],  # DOCX has no fixed page boundaries
        paragraphs=paragraphs,
        tables=tables,
        full_text=full_text,
        character_count=char_count,
        page_count=1,
        extraction_method="python-docx",
        warnings=warnings,
    )
    return DocumentReadResult(ok=True, document=content, warnings=warnings)


def read_txt(
    data: bytes,
    filename: str = "",
    media_type: str = "",
    *,
    max_chars: int = DEFAULT_MAX_CHARS,
) -> DocumentReadResult:
    """Read a TXT document supporting UTF-8 and legacy encodings with line preservation."""
    source = DocumentSource(
        filename=filename,
        media_type=media_type or "text/plain",
        extension=".txt",
        source_type="txt",
    )
    warnings: list[str] = []

    if not data:
        return DocumentReadResult(
            ok=False,
            error_code="EMPTY_FILE",
            error_message="The uploaded text file is empty.",
            warnings=warnings,
        )

    text: str = ""
    encoding_used = "utf-8"
    if isinstance(data, str):
        text = data
    else:
        try:
            text = data.decode("utf-8-sig")  # handles BOM if present
        except UnicodeDecodeError:
            try:
                text = data.decode("latin-1")
                encoding_used = "latin-1"
                warnings.append("File decoded using fallback latin-1 encoding.")
            except Exception as exc:
                return DocumentReadResult(
                    ok=False,
                    error_code="MALFORMED_DOCUMENT",
                    error_message=f"Could not decode text file: {exc}",
                    warnings=warnings,
                )


    text = text.replace("\r\n", "\n").replace("\r", "\n")
    if len(text) > max_chars:
        text = text[:max_chars]
        warnings.append(f"Text truncated at maximum limit of {max_chars} characters.")

    full_text = text.strip()
    char_count = len(full_text)

    if char_count == 0:
        return DocumentReadResult(
            ok=False,
            error_code="EMPTY_FILE",
            error_message="The uploaded text file is empty or contains only whitespace.",
            warnings=warnings,
        )

    paragraphs: list[DocumentParagraph] = []
    for idx, line in enumerate(text.splitlines()):
        line_clean = line.strip()
        if line_clean:
            paragraphs.append(
                DocumentParagraph(
                    text=line_clean,
                    style="Normal",
                    paragraph_index=idx,
                )
            )

    content = DocumentContent(
        source=source,
        pages=[DocumentPage(page_number=1, text=full_text, character_count=char_count)],
        paragraphs=paragraphs,
        tables=[],
        full_text=full_text,
        character_count=char_count,
        page_count=1,
        extraction_method=f"plain_text ({encoding_used})",
        warnings=warnings,
    )
    return DocumentReadResult(ok=True, document=content, warnings=warnings)


def read_document(
    data: bytes,
    filename: str = "",
    media_type: str = "",
    *,
    max_bytes: int = DEFAULT_MAX_BYTES,
    max_chars: int = DEFAULT_MAX_CHARS,
) -> DocumentReadResult:
    """Unified document reader dispatching to PDF, DOCX, or TXT extractors with safety checks."""
    if not data or len(data) == 0:
        return DocumentReadResult(
            ok=False,
            error_code="EMPTY_FILE",
            error_message="The uploaded file is empty.",
        )

    if len(data) > max_bytes:
        return DocumentReadResult(
            ok=False,
            error_code="FILE_TOO_LARGE",
            error_message=f"File size ({len(data)} bytes) exceeds the maximum allowed limit of {max_bytes} bytes.",
        )

    fmt = detect_format(data, filename=filename, media_type=media_type)

    if fmt == "pdf":
        return read_pdf(data, filename=filename, media_type=media_type, max_chars=max_chars)
    elif fmt == "docx":
        return read_docx(data, filename=filename, media_type=media_type, max_chars=max_chars)
    elif fmt == "txt":
        return read_txt(data, filename=filename, media_type=media_type, max_chars=max_chars)
    else:
        ext = os.path.splitext(filename)[1] if filename else ""
        return DocumentReadResult(
            ok=False,
            error_code="UNSUPPORTED_FORMAT",
            error_message=(
                f"Unsupported document format: '{ext or media_type or 'unknown'}'. "
                "Supported formats are PDF (.pdf), Word (.docx), and Plain Text (.txt)."
            ),
        )
